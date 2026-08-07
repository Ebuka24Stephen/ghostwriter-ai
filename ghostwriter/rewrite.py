import io
import re

from docx import Document

from . import config, llm


def is_heading(text):
    stripped = text.strip()
    if not stripped:
        return True
    if stripped.startswith("#") and len(stripped) < 60:
        return True
    if len(stripped) < 60 and re.match(r"^[A-Z][A-Z\s\d\.]+$", stripped):
        return True
    if len(stripped) < 60 and re.match(r"^CHAPTER\s", stripped, re.IGNORECASE):
        return True
    if len(stripped) < 60 and re.match(r"^\d+\.\d+", stripped):
        return True
    return False


def split_document(text):
    paragraphs = []
    for p in re.split(r"\n\n|(?=###\s)", text):
        p = p.strip()
        if len(p) > 10:
            paragraphs.append(p)
    return paragraphs


def extract_heading_prefix(text):
    match = re.match(r"^(###\s+\d+\.\d+\s+[A-Z][A-Z\s]+?)\s+[A-Z][a-z]", text)
    if match:
        return match.group(1).strip()
    return None


KEEP_RATIO = config.KEEP_RATIO
CHUNK_SIZE = config.REWRITE_CHUNK_SIZE

DUP_THRESHOLD = 0.85


def _normalize(s):
    return re.sub(r"[^a-z0-9 ]", " ", s.lower())


def is_duplicate(candidate, seen):
    cand = set(_normalize(candidate).split())
    if len(cand) < 4:
        return False
    for s in seen:
        seen_tokens = set(_normalize(s).split())
        intersection = cand & seen_tokens
        overlap_cand = len(intersection) / max(1, len(cand))
        overlap_seen = len(intersection) / max(1, len(seen_tokens))
        if overlap_cand >= DUP_THRESHOLD or overlap_seen >= DUP_THRESHOLD:
            return True
    return False


def _overlap(a, b):
    ta = set(_normalize(a).split())
    tb = set(_normalize(b).split())
    if not ta:
        return 0.0
    return len(ta & tb) / len(ta)


def _para_is_dup(text, seen):
    cand = set(_normalize(text).split())
    if len(cand) < 10:
        return False
    for s in seen:
        seent = set(_normalize(s).split())
        inter = cand & seent
        if inter and len(inter) / len(cand) >= DUP_THRESHOLD:
            return True
    return False


def split_sentences(text):
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"'`])", text) if s.strip()]


def plan_keep_indices(n_sentences, keep_ratio=KEEP_RATIO):
    if n_sentences < 2 or keep_ratio <= 0:
        return set()
    n_keep = max(1, round(n_sentences * keep_ratio))
    step = n_sentences / n_keep
    return {min(n_sentences - 1, round(k * step)) for k in range(n_keep)}


def _parse_numbered(text):
    result = {}
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", lines[i])
        if m:
            num = int(m.group(1))
            parts = [m.group(2)]
            i += 1
            while i < len(lines):
                if re.match(r"^\s*\d+[.)]\s+", lines[i]):
                    break
                parts.append(lines[i])
                i += 1
            result[num] = " ".join(parts).strip()
        else:
            i += 1
    return result


MAX_GROWTH = 1.3


def cap_paragraph(text, source_text):
    limit = int(len(source_text.split()) * MAX_GROWTH)
    if len(text.split()) <= limit:
        return text
    source_sentences = set(split_sentences(source_text))
    capped = []
    total = 0
    stopped = False
    for s in split_sentences(text):
        n = len(s.split())
        if s in source_sentences:
            capped.append(s)
            total += n
            continue
        if stopped:
            continue
        if capped and total + n > limit:
            stopped = True
            continue
        capped.append(s)
        total += n
    if not capped:
        return text
    return " ".join(capped)


def cap_document(input_doc, output_text):
    in_paras = input_doc.split("\n\n")
    out_paras = output_text.split("\n\n")
    capped = []
    for i, op in enumerate(out_paras):
        if i < len(in_paras):
            capped.append(cap_paragraph(op, in_paras[i]))
        else:
            capped.append(op)
    return normalize_spacing("\n\n".join(capped))


_TAIL_PATTERNS = [
    re.compile(r",\s+and this is\s+.*?\.?$", re.IGNORECASE),
    re.compile(r",\s+and it is\s+.*?\.?$", re.IGNORECASE),
    re.compile(r",\s+thus\s+[a-z]+ing\s+.*?\.?$", re.IGNORECASE),
    re.compile(r",\s+which is\s+(?:very|an|a|also|integral)\s+.*?\.?$", re.IGNORECASE),
    re.compile(
        r",\s+and\s+.{3,90}?\b(?:inadequate|in need of|capable|very much|"
        r"crucial|important|significant|essential|vital|beneficial|worth|"
        r"necessary|limitation|challenge|problem|often|of such|effective|"
        r"useful|valuable|concerning|problematic|promising)\b.*?\.?$",
        re.IGNORECASE,
    ),
]

_FOOTER_PATTERNS = [
    re.compile(r"^\s*Page\s+\d+\s+of\s+\d+\s+.*$", re.IGNORECASE),
    re.compile(r"^\s*Submission\s+ID\s+trn:oid\s*.*$", re.IGNORECASE),
    re.compile(r"^\s*AI\s+Writing\s+Submission.*$", re.IGNORECASE),
]


def sanitize_footers(text):
    out = []
    for line in text.splitlines():
        if any(p.match(line) for p in _FOOTER_PATTERNS):
            continue
        out.append(line)
    return "\n".join(out)


def normalize_spacing(text):
    return re.sub(r"(\S)\s+([,;:)])", r"\1\2", text)


def strip_added_tail(rewritten, original):
    orig_lower = original.lower()
    cut = None
    for pattern in _TAIL_PATTERNS:
        m = pattern.search(rewritten)
        if m:
            tail = m.group(0)
            if tail.lower() in orig_lower:
                continue
            if cut is None or m.start() < cut:
                cut = m.start()
    if cut is None:
        return rewritten
    candidate = rewritten[:cut].rstrip(" ,")
    if len(candidate.split()) < 3:
        return rewritten
    return candidate


def rewrite_document(document, system_prompt):
    document = sanitize_footers(document)
    paras = [p for p in document.split("\n\n") if p.strip()]
    deduped = []
    seen_paras = []
    for para in paras:
        stripped = para.strip()
        if is_heading(stripped) or not _para_is_dup(stripped, seen_paras):
            deduped.append(stripped)
            if not is_heading(stripped):
                seen_paras.append(stripped)
    document = "\n\n".join(deduped)

    seen_kept = []
    to_rewrite = []
    plan = []
    counter = 0

    for para in document.split("\n\n"):
        stripped = para.strip()
        if not stripped:
            plan.append([])
            continue
        if is_heading(stripped):
            plan.append([stripped])
            continue
        sentences = split_sentences(stripped)
        if len(sentences) < 2:
            plan.append(sentences)
            continue
        keep_indices = plan_keep_indices(len(sentences))
        para_plan = []
        for i, s in enumerate(sentences):
            if i in keep_indices and not is_duplicate(s, seen_kept):
                seen_kept.append(s)
                para_plan.append(s)
            else:
                counter += 1
                para_plan.append("REWRITE:%d" % counter)
                prev_s = sentences[i - 1] if i > 0 else None
                next_s = sentences[i + 1] if i + 1 < len(sentences) else None
                to_rewrite.append((counter, s, prev_s, next_s))
        plan.append(para_plan)

    rewrites = {}
    original_map = {n: s for n, s, _, _ in to_rewrite}
    for start in range(0, len(to_rewrite), CHUNK_SIZE):
        chunk = to_rewrite[start:start + CHUNK_SIZE]
        lines = []
        for n, s, prev_s, next_s in chunk:
            if prev_s and next_s:
                ctx = '\n   (Between: "%s" ... "%s")' % (prev_s, next_s)
            elif prev_s:
                ctx = '\n   (Previous sentence: "%s")' % prev_s
            elif next_s:
                ctx = '\n   (Next sentence: "%s")' % next_s
            else:
                ctx = ""
            lines.append("%d. %s%s" % (n, s, ctx))
        numbered = "\n".join(lines)
        prompt = f"""Rewrite each numbered sentence below in the author's voice as described in the system prompt.

{numbered}

Rules:
- Produce exactly ONE rewritten sentence per number, in the same order
- Keep the same numbering
- Do not add, remove, merge, or split any sentences
- Say the same idea in clearly DIFFERENT words and sentence structure. This must NOT be a near-copy of the original — change at least half the words and reorder the clauses
- Do NOT add new facts, examples, or ideas
- Do NOT append any new clause that adds evaluation, commentary, or a conclusion (no "and this is...", "thus ...ing", "which is...", or similar)
- Keep all facts, numbers, and any citation exactly as written inside its parentheses, e.g. "(Ismail et al., 2023)"
- Match the original's length roughly, not exactly — the goal is different wording, not different size
- Use the "(Between: ...)" context only to fit the sentence naturally; never copy words from it into your rewritten sentence
- Self-check before answering: compare each sentence you wrote against the author's example passages in the system prompt. If it sounds too clean, smooth, modern, or AI-like, rewrite it again internally. Also confirm it is NOT a near-copy of the original sentence
- Return only the numbered rewritten sentences, nothing else"""

        result = llm.ask(prompt, system_prompt=system_prompt, temperature=config.REWRITE_TEMPERATURE)
        rewrites.update(_parse_numbered(result))

    for attempt in range(2):
        too_similar = [
            n for n, orig in original_map.items()
            if rewrites.get(n) and _overlap(rewrites[n], orig) >= 0.7
        ]
        if not too_similar:
            break
        lines = "\n".join(
            "%d. %s" % (n, original_map[n]) for n in too_similar
        )
        prompt = f"""Each numbered sentence below was rewritten but the result is still too close to the original — it looks copied. Rewrite it again with clearly different wording and sentence order, keeping the same meaning and any citation.

{lines}

Rules:
- Change at least half the words; reorder the clauses; do not keep the original's phrasing
- Keep all facts and any citation exactly
- Do not add new ideas or append evaluative clauses
- Return only the numbered rewritten sentences, nothing else"""
        result = llm.ask(prompt, system_prompt=system_prompt, temperature=config.REWRITE_TEMPERATURE)
        rewrites.update(_parse_numbered(result))

    rebuilt = []
    for para_plan in plan:
        parts = []
        for item in para_plan:
            if isinstance(item, str) and item.startswith("REWRITE:"):
                n = int(item.split(":", 1)[1])
                rewritten = rewrites.get(n, original_map.get(n, ""))
                parts.append(strip_added_tail(rewritten, original_map.get(n, "")))
            else:
                parts.append(item)
        rebuilt.append(" ".join(p for p in parts if p))

    return cap_document(document, "\n\n".join(rebuilt))


def rewrite_docx(file_bytes: bytes, system_prompt: str) -> bytes:
    doc = Document(io.BytesIO(file_bytes))
    original = list(doc.paragraphs)

    body = []
    for i, p in enumerate(original):
        text = p.text.strip()
        if not text:
            continue
        if is_heading(text):
            continue
        if len(text) < 20:
            continue
        body.append((text, i))

    seen = []
    filtered = []
    for text, i in body:
        if not _para_is_dup(text, seen):
            filtered.append((text, i))
            seen.append(text)
    body_paras = [t for t, _ in filtered]
    body_indices = {i for _, i in filtered}

    rewritten_body = rewrite_document("\n\n".join(body_paras), system_prompt)
    rewritten_paras = [p.strip() for p in rewritten_body.split("\n\n") if p.strip()]

    out = Document()
    idx = 0
    for i, p in enumerate(original):
        text = p.text.strip()
        if not text:
            continue

        if i in body_indices:
            rewritten = rewritten_paras[idx] if idx < len(rewritten_paras) else text
            out.add_paragraph(cap_paragraph(rewritten, text))
            idx += 1
        else:
            out.add_paragraph(text, style=p.style.name if p.style.name.startswith("Heading") else None)

    buf = io.BytesIO()
    out.save(buf)
    buf.seek(0)
    return buf.getvalue()
